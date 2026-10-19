"""Phase 1 — ENG-05 robust parsing + SEC-08 hardened file pipeline."""
import io
import zipfile

import pytest

from studybuddy.parsing import (
    ParseError,
    detect_fields,
    detect_type,
    parse_document,
    validate_upload,
)


def make_pdf(text_lines):
    """Minimal valid PDF with the given text lines."""
    content = b"BT /F1 12 Tf 72 740 Td 14 TL\n"
    for line in text_lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content += b"(" + safe.encode("latin-1", "ignore") + b") Tj T*\n"
    content += b"ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(pdf)
    pdf += b"xref\n0 %d\n" % (len(objects) + 1)
    pdf += b"0000000000 65535 f \n"
    for off in offsets:
        pdf += b"%010d 00000 n \n" % off
    pdf += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (
        len(objects) + 1, xref)
    return pdf


def make_docx(paragraphs, tables=()):
    from docx import Document
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    for rows in tables:
        tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                tbl.cell(i, j).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


RESUME_LINES = [
    "Aarav Sharma",
    "aarav.sharma@example.com | +91-9876543210",
    "https://github.com/aaravsharma | https://linkedin.com/in/aaravsharma",
    "SUMMARY",
    "Computer Science undergraduate with strong Python skills.",
    "EDUCATION",
    "B.Tech Computer Science, Graphic Era Hill University, 2023 - 2027",
    "EXPERIENCE",
    "Machine Learning Intern, 6 months building Flask APIs.",
    "SKILLS",
    "Python, Java, SQL, Machine Learning, Flask, Git",
]


class TestContentValidation:
    def test_pdf_detected_by_magic(self):
        assert detect_type(make_pdf(["hi"])) == "pdf"

    def test_docx_detected_by_magic(self):
        assert detect_type(make_docx(["hi"])) == "docx"

    def test_extension_spoofing_rejected(self):
        # EXE bytes renamed to .pdf -> rejected (SEC-08.1)
        with pytest.raises(ParseError):
            validate_upload(b"MZ\x90\x00 evil", filename="resume.pdf")

    def test_text_disguised_as_pdf_rejected(self):
        with pytest.raises(ParseError):
            validate_upload(b"just some text", filename="resume.pdf")

    def test_legacy_doc_rejected_with_helpful_message(self):
        with pytest.raises(ParseError, match="Legacy .doc"):
            validate_upload(b"\xd0\xcf\x11\xe0" + b"\x00" * 20, filename="old.doc")

    def test_oversized_rejected(self):
        big = b"%PDF-" + b"\x00" * (6 * 1024 * 1024)
        with pytest.raises(ParseError, match="too large"):
            validate_upload(big, filename="big.pdf", kind="resume")

    def test_zip_bomb_rejected(self):
        # DOCX that decompresses beyond the cap
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", b"<Types/>")
            zf.writestr("word/document.xml",
                        b"<w:doc>" + b"A" * (60 * 1024 * 1024) + b"</w:doc>")
        with pytest.raises(ParseError, match="unexpectedly large|corrupted|recognize"):
            validate_upload(buf.getvalue(), filename="bomb.docx")


class TestPdfParsing:
    def test_extracts_text_and_fields(self):
        res = parse_document(make_pdf(RESUME_LINES), filename="r.pdf")
        assert res.ok
        assert "Aarav Sharma" in res.text
        assert res.fields["email"] == "aarav.sharma@example.com"
        assert "9876543210" in res.fields["phone"]
        assert res.fields["name"] == "Aarav Sharma"
        assert any("github.com" in u for u in res.fields["links"])
        assert "EDUCATION" in " ".join(res.fields["sections"]).upper()

    def test_layout_and_quality_present(self):
        res = parse_document(make_pdf(RESUME_LINES), filename="r.pdf")
        assert res.pages == 1
        assert "extractor" in res.layout
        assert res.quality["chars_per_page"] > 50
        assert res.quality["garbled_ratio"] < 0.05

    def test_garbage_pdf_gives_friendly_error(self):
        with pytest.raises(ParseError, match="Could not read|Could not extract"):
            parse_document(b"%PDF-1.4\n%garbage" + b"\x00" * 100,
                           filename="bad.pdf")

    def test_too_many_pages_rejected(self):
        pdf = make_pdf(["x"])
        # fake a 7-page resume by concatenating catalogs is complex; instead
        # directly exercise the page-limit check via a patched count
        import studybuddy.parsing as P
        orig = P._pdf_page_count
        P._pdf_page_count = lambda data: 99
        try:
            with pytest.raises(ParseError, match="99 pages"):
                P.extract_pdf(pdf, max_pages=6)
        finally:
            P._pdf_page_count = orig


class TestDocxParsing:
    def test_paragraphs_and_tables(self):
        data = make_docx(
            ["Aarav Sharma", "aarav.sharma@example.com", "SKILLS"],
            tables=[[("Python", "Expert"), ("SQL", "Intermediate")]],
        )
        res = parse_document(data, filename="r.docx")
        assert res.ok
        assert "Aarav Sharma" in res.text
        assert "Python | Expert" in res.text
        assert res.layout["has_tables"] is True
        assert any("table" in w for w in res.warnings)
        assert res.fields["email"] == "aarav.sharma@example.com"

    def test_corrupt_docx_rejected(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("[Content_Types].xml", b"<Types/>")
            zf.writestr("word/document.xml", b"<broken")
        # passes structural validation, fails at parse time with friendly error
        with pytest.raises(ParseError):
            parse_document(buf.getvalue(), filename="c.docx")


class TestFieldDetection:
    def test_empty_text_gives_empty_fields(self):
        assert detect_fields("") == {}

    def test_name_heuristic_skips_contact_lines(self):
        fields = detect_fields("aarav@example.com\n+91-9876543210\nAarav Sharma\nEngineer")
        assert fields["name"] == "Aarav Sharma"

    def test_sections_detected(self):
        fields = detect_fields("EDUCATION\nB.Tech\nEXPERIENCE\nIntern")
        assert len(fields["sections"]) == 2


class TestParsePreviewEndpoint:
    def test_preview_returns_fields(self, client):
        data = {"file": (io.BytesIO(make_pdf(RESUME_LINES)), "r.pdf")}
        r = client.post("/api/parse_preview", data=data,
                        content_type="multipart/form-data")
        assert r.status_code == 200
        body = r.get_json()
        assert body["fields"]["email"] == "aarav.sharma@example.com"
        assert "Aarav Sharma" in body["text"]
        assert "quality" in body and "warnings" in body

    def test_preview_rejects_spoofed_file(self, client):
        data = {"file": (io.BytesIO(b"MZ evil"), "r.pdf")}
        r = client.post("/api/parse_preview", data=data,
                        content_type="multipart/form-data")
        assert r.status_code == 400
